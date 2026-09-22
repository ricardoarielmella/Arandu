# -*- coding: utf-8 -*-
import openpyxl, json, re, unicodedata
F="/root/.claude/uploads/7bf9e701-0a19-5a40-9ae0-57a559b58c46/51ca6652-Copia_de_ROL_DE_COMBATE_ARANDU_2026_INTEGRAL_ACTUALIZADA_22_SEP.xlsx"
wb=openpyxl.load_workbook(F, data_only=True)

def t(c):
    if c is None: return ""
    if isinstance(c,float) and c==int(c): c=int(c)
    return re.sub(r"\s+"," ",str(c)).strip()

ARMAS=["FUSIL FAL","FUSIL 5,56 mm","Pistola 9 mm","Pistola Amet 9 mm","Fusil .50",
       "Fusil Tir Esp","Amet 7,62 mm","Amet 5,56 mm","AT 4","CARL GUSTAV"]

filas=[]; seccion=""
ws=wb['ROL DE COMBATE ARANDÚ']
for i,r in enumerate(ws.iter_rows(min_row=9, values_only=True),9):
    v=[t(c) for c in r]
    if not any(v): continue
    if v[0] and not v[4] and not v[5]:
        seccion=v[0]; continue
    if not v[4] and not v[5]: continue
    filas.append({"hoja":"rol","fila":i,"seccion":seccion,"ord":v[0],"nroElem":v[1],
        "grado":v[2],"arma":v[3],"nombre":v[4],"dni":v[5],"rol":v[6],"destino":v[7],
        "armas":{ARMAS[k]:v[8+k] for k in range(10) if v[8+k]},
        "veh":v[18],"lic":v[19],"licVto":v[20]})

ws2=wb['CONDUCTORES']; seccion=""
for i,r in enumerate(ws2.iter_rows(min_row=7, values_only=True),7):
    v=[t(c) for c in r]+[""]*12
    if not any(v[:12]): continue
    if v[0] and not v[4] and not v[5]:
        seccion=v[0]; continue
    if v[0]=="NRO Ord": continue
    if not v[4] and not v[5]: continue
    filas.append({"hoja":"cond","fila":i,"seccion":seccion,"ord":v[0],"nroElem":v[1],
        "grado":v[2],"arma":v[3],"nombre":v[4],"dni":v[5],"rol":v[6],"destino":v[7],
        "armas":{}, "licVto":v[8], "veh":v[9], "categoria":v[10], "obs":v[11], "lic":""})

json.dump(filas, open("rol_filas.json","w"), ensure_ascii=False, indent=1)
print("filas:", len(filas), "| rol:", sum(1 for f in filas if f['hoja']=='rol'),
      "| cond:", sum(1 for f in filas if f['hoja']=='cond'))
print("con armamento:", sum(1 for f in filas if f['armas']))
print("con vehiculo:", sum(1 for f in filas if f['veh']))
print("con licencia:", sum(1 for f in filas if f['lic'] or f['licVto']))
print("con rol:", sum(1 for f in filas if f['rol']))
print("con destino:", sum(1 for f in filas if f['destino']))
print("secciones:", len({f['seccion'] for f in filas}))
